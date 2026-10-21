#!/usr/bin/env python3
"""REPORT.md -> REPORT.docx (python-docx) -> REPORT.pdf (LibreOffice).

Handles the Markdown subset used in REPORT.md: headings, paragraphs, **bold**, *italic*,
`code`, bullet/numbered lists, tables, images, fenced code, <!-- pagebreak --> and <!-- toc -->.
"""
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent


def add_runs(p, text, size=None):
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text.replace("&nbsp;", " ")):
        if not part:
            continue
        if part.startswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*"):
            r = p.add_run(part[1:-1]); r.italic = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
        else:
            r = p.add_run(part)
        if size:
            r.font.size = Pt(size)


def toc(doc):
    p = doc.add_paragraph()
    r = p.add_run()
    for tag, txt in (("begin", None), (None, 'TOC \\o "1-3" \\h \\z \\u'), ("separate", None), (None, None), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag); r._r.append(el)
        elif txt:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt; r._r.append(el)
        else:
            t = OxmlElement("w:t"); t.text = "Right-click → Update Field to refresh the table of contents."; r._r.append(t)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)


def build():
    lines = (HERE / "REPORT.md").read_text().splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = Cm(3.0); sec.right_margin = Cm(2.0); sec.top_margin = sec.bottom_margin = Cm(2.2)
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.3
    for lvl, size in ((1, 16), (2, 15), (3, 13)):
        hs = doc.styles[f"Heading {lvl}"]; hs.font.name = "Times New Roman"; hs.font.size = Pt(size)
        hs.font.color.rgb = RGBColor(0, 0, 0); hs.font.bold = True
        hs.element.rPr.rFonts.set(qn("w:asciiTheme"), "") if False else None
    # page numbers in footer
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run()
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag); r._r.append(el)
        else:
            el = OxmlElement("w:instrText"); el.text = txt; r._r.append(el)

    cover = True
    i = 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if s == "<!-- pagebreak -->":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE); cover = False; i += 1; continue
        if s == "<!-- toc -->":
            toc(doc); i += 1; continue
        if s.startswith("```"):
            buf = []; i += 1
            while not lines[i].strip().startswith("```"):
                buf.append(lines[i]); i += 1
            p = doc.add_paragraph(); r = p.add_run("\n".join(buf)); r.font.name = "Consolas"; r.font.size = Pt(9)
            p.paragraph_format.line_spacing = 1.0
            i += 1; continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if all(not c for c in rows[0]):        # header-less signature table
                rows = rows[1:]
                t = doc.add_table(rows=len(rows), cols=len(rows[0]))
                hdr = False
            else:
                t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"; hdr = True
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, row in enumerate(rows):
                for ci, c in enumerate(row):
                    cell = t.cell(ri, ci); cell.text = ""
                    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
                    add_runs(p, f"**{c}**" if hdr and ri == 0 and c and not c.startswith("**") else c, size=10)
                    if hdr and ri == 0:
                        shade(cell, "D9E2F3")
            doc.add_paragraph()
            continue
        m = re.match(r"!\[(.*)\]\((.*)\)", s)
        if m:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(HERE / m.group(2)), width=Cm(15.5))
            c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = c.add_run(m.group(1)); r.italic = True; r.font.size = Pt(10)
            i += 1; continue
        m = re.match(r"(#+)\s+(.*)", s)
        if m:
            lvl = len(m.group(1))
            if cover and lvl == 1:
                for _ in range(4):
                    doc.add_paragraph()
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(m.group(2)); r.bold = True; r.font.size = Pt(20)
            else:
                h = doc.add_heading(m.group(2), level=min(lvl, 3) if lvl > 1 else 1)
                if lvl == 2:
                    h.style = doc.styles["Heading 1"]
                elif lvl == 3:
                    h.style = doc.styles["Heading 2"]
            i += 1; continue
        m = re.match(r"(\d+)\.\s+(.*)", s) or re.match(r"[-*]\s+(.*)", s)
        if m:
            if s[0].isdigit():
                p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.8)
                p.paragraph_format.first_line_indent = Cm(-0.6)
                add_runs(p, f"{m.group(1)}. {m.group(2)}")
            else:
                p = doc.add_paragraph(style="List Bullet"); add_runs(p, m.groups()[-1])
            i += 1; continue
        if s:
            p = doc.add_paragraph()
            if cover:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, s, size=14 if cover else None)
        i += 1
    out = HERE / "REPORT.docx"
    doc.save(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(HERE), str(out)],
                   check=True, stdout=subprocess.DEVNULL)
    print("wrote", out, HERE / "REPORT.pdf")


if __name__ == "__main__":
    build()
